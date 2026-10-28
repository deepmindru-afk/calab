#!/usr/bin/env python3
"""Blank templates of «Создать» in the Disk (ADR-0143): apps/server/internal/office/templates/
blank.{docx,xlsx,pptx}, embedded by the server (office/create.go), which sets the text language
(en-US here) and the sheet name per request.

  python3 -m venv /tmp/of && /tmp/of/bin/pip install openpyxl python-docx python-pptx
  /tmp/of/bin/python -I tools/office-fidelity/templates.py apps/server/internal/office/templates

The files come from the libraries' own Office-made templates, then: A4 with 2 cm margins (docx),
16:9 with the layouts stretched to it (pptx) and one empty title slide, neutral document
properties (no author, a fixed date), and a deterministic zip (fixed entry times), so a re-run
gives the same bytes.
"""
import datetime as dt
import io
import re
import sys
import zipfile
from pathlib import Path

FIXED = dt.datetime(2026, 1, 1, 0, 0, 0)


def neutral(props) -> None:
    for k in ("title", "subject", "keywords", "category"):
        setattr(props, k, "")
    for k in ("author", "last_modified_by", "comments", "creator", "lastModifiedBy", "description"):
        if hasattr(props, k):
            try:
                setattr(props, k, "")
            except (AttributeError, TypeError, ValueError):
                pass
    props.created = FIXED
    props.modified = FIXED


def stable_zip(raw: bytes) -> bytes:
    """The same entries in the same order with a fixed time: reproducible bytes. The libraries'
    template thumbnail (a picture of THEIR sample page) is dropped with its relationship."""
    src = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            if info.filename == "docProps/thumbnail.jpeg":
                continue
            data = src.read(info.filename)
            if info.filename == "_rels/.rels":
                data = re.sub(rb'<Relationship [^>]*Target="docProps/thumbnail.jpeg"[^>]*/>', b"", data)
            if info.filename == "docProps/core.xml":  # openpyxl stamps the save time
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", data)
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)
    return out.getvalue()


def docx() -> bytes:
    from docx import Document
    from docx.shared import Cm

    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2)
    neutral(d.core_properties)
    d.core_properties.revision = 1
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Sheet1"
    neutral(wb.properties)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    k = 16 / 12  # 4:3 → 16:9 at the same height
    prs.slide_width = Emu(12192000)
    for shapes in [prs.slide_master.shapes] + [layout.shapes for layout in prs.slide_layouts]:
        for sh in shapes:
            if sh.left is not None and sh.width is not None:
                sh.left, sh.width = Emu(int(sh.left * k)), Emu(int(sh.width * k))
    prs.slides.add_slide(prs.slide_layouts[0])
    neutral(prs.core_properties)
    prs.core_properties.revision = 1
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "apps/server/internal/office/templates")
    out.mkdir(parents=True, exist_ok=True)
    for name, make in (("blank.docx", docx), ("blank.xlsx", xlsx), ("blank.pptx", pptx)):
        (out / name).write_bytes(stable_zip(make()))
        print(out / name, (out / name).stat().st_size)
