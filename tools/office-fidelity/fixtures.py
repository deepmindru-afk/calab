#!/usr/bin/env python3
"""Office fidelity fixtures (ADR-0143): realistic files like the Disk files bots read.

  python3 -m venv /tmp/of && /tmp/of/bin/pip install openpyxl python-docx python-pptx
  /tmp/of/bin/python tools/office-fidelity/fixtures.py OUT_DIR

Writes OUT_DIR/vacations.xlsx («График отпусков»-like), OUT_DIR/contract.docx («Лицензионный
договор»-like), OUT_DIR/deck.pptx (layouts Title, Title and Content with levels and notes, Two Content,
Title Only with a table, a column chart and a styled text box, 16:9) and OUT_DIR/staff.csv (UTF-8, «;», quoted fields with «;» and
newlines). The legacy doc / xls / ppt are made from these by the Document Server itself in
bench.mjs (no Python library writes them). Every feature compare.py checks is in them: formulas (same-sheet, cross-sheet,
absolute refs, functions with Russian text), merged cells, date / number / percent / currency
formats, column widths and row heights, frozen panes, conditional formatting, data validation, a
defined name, styles (fonts, fills, borders, alignment); a docx with heading and list styles, a
styled table with a header row and merged cells, page headers and footers (different first
page), page margins and orientation, a page break, bold / italic runs.
"""
import datetime as dt
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation


def xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "График"
    staff = wb.create_sheet("Сотрудники")
    totals = wb.create_sheet("Итоги")

    thin = Side(style="thin", color="999999")
    head_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    head_fill = PatternFill("solid", fgColor="305496")

    # Сотрудники: names, departments, salaries (currency), rates (percent).
    staff.append(["Таб. №", "ФИО", "Отдел", "Оклад", "Премия, %", "Дата приёма"])
    people = [
        (101, "Анна Иванова", "Разработка", 185000, 0.15, dt.date(2021, 3, 1)),
        (102, "Борис Петров", "Продажи", 140000, 0.25, dt.date(2022, 6, 15)),
        (103, "Вера Смирнова", "HR", 120000, 0.10, dt.date(2019, 11, 20)),
        (104, "Григорий Орлов", "Разработка", 210000, 0.2, dt.date(2020, 1, 9)),
    ]
    for p in people:
        staff.append(list(p))
    for c in staff[1]:
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in range(2, 2 + len(people)):
        staff[f"D{r}"].number_format = '#,##0.00 "₽"'
        staff[f"E{r}"].number_format = "0%"
        staff[f"F{r}"].number_format = "DD.MM.YYYY"
    for col, w in zip("ABCDEF", (8, 24, 14, 14, 11, 13)):
        staff.column_dimensions[col].width = w
    staff.freeze_panes = "B2"
    wb.defined_names["Оклады"] = DefinedName("Оклады", attr_text="Сотрудники!$D$2:$D$5")

    # График: a merged title, dates, day counts by formula, cross-sheet lookups.
    ws.merge_cells("A1:F1")
    ws["A1"] = "График отпусков 2026"
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 28
    ws.append([])
    ws.append(["Таб. №", "ФИО", "Начало", "Конец", "Дней", "Отдел"])
    for c in ws[3]:
        c.font, c.fill = head_font, head_fill
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    rows = [(101, dt.date(2026, 7, 1), dt.date(2026, 7, 14)), (102, dt.date(2026, 7, 15), dt.date(2026, 7, 28)),
            (103, dt.date(2026, 8, 3), dt.date(2026, 8, 16)), (104, dt.date(2026, 12, 21), dt.date(2027, 1, 9))]
    for i, (tab, start, end) in enumerate(rows, start=4):
        ws[f"A{i}"] = tab
        ws[f"B{i}"] = f'=VLOOKUP(A{i},Сотрудники!$A$2:$F$5,2,FALSE)'
        ws[f"C{i}"], ws[f"D{i}"] = start, end
        ws[f"C{i}"].number_format = ws[f"D{i}"].number_format = "DD.MM.YYYY"
        ws[f"E{i}"] = f"=D{i}-C{i}+1"
        ws[f"F{i}"] = f'=INDEX(Сотрудники!$C$2:$C$5,MATCH(A{i},Сотрудники!$A$2:$A$5,0))'
        for col in "ABCDEF":
            ws[f"{col}{i}"].border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws["A9"] = "Итого дней"
    ws.merge_cells("A9:D9")
    ws["A9"].alignment = Alignment(horizontal="right")
    ws["E9"] = "=SUM(E4:E7)"
    ws["E9"].font = Font(bold=True)
    for col, w in zip("ABCDEF", (8, 22, 12, 12, 8, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    ws.conditional_formatting.add("E4:E7", CellIsRule(operator="greaterThan", formula=["14"], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(color="9C0006")))
    ws.conditional_formatting.add("C4:C7", FormulaRule(formula=["MONTH(C4)=12"], fill=PatternFill("solid", fgColor="FFEB9C")))
    dv = DataValidation(type="list", formula1='"Разработка,Продажи,HR"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("F4:F7")

    # Итоги: functions over a named range and the other sheets, percent and number formats.
    totals["A1"], totals["B1"] = "Фонд окладов", "=SUM(Оклады)"
    totals["A2"], totals["B2"] = "Средний оклад", "=AVERAGE(Сотрудники!D2:D5)"
    totals["A3"], totals["B3"] = "Премии", "=SUMPRODUCT(Сотрудники!D2:D5,Сотрудники!E2:E5)"
    totals["A4"], totals["B4"] = "Доля премий", "=B3/B1"
    totals["A5"], totals["B5"] = "Дней отпуска", "=График!E9"
    totals["A6"], totals["B6"] = "Разработчиков", '=COUNTIF(Сотрудники!C2:C5,"Разработка")'
    totals["A7"], totals["B7"] = "Отчёт", '=TEXT(DATE(2026,12,31),"DD.MM.YYYY")&" — "&B6&" чел."'
    for r in (1, 2, 3):
        totals[f"B{r}"].number_format = '#,##0 "₽"'
    totals["B4"].number_format = "0.0%"
    totals.column_dimensions["A"].width = 20
    totals.column_dimensions["B"].width = 26
    totals.conditional_formatting.add("B1:B3", ColorScaleRule(start_type="min", start_color="F8696B", end_type="max", end_color="63BE7B"))
    wb.save(path)


def docx(path: Path) -> None:
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.different_first_page_header_footer = True
    sec.first_page_header.paragraphs[0].text = "ООО «Калаба» · ШАБЛОН"
    sec.header.paragraphs[0].text = "Лицензионный договор № ___ · стр."
    sec.footer.paragraphs[0].text = "Лицензиар ________   Лицензиат ________"
    sec.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)

    t = d.add_paragraph("ЛИЦЕНЗИОННЫЙ ДОГОВОР № ___", style="Title")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = d.add_paragraph("г. Москва")
    p.add_run("\t«___» ________ 2026 г.").italic = True
    d.add_heading("1. Предмет договора", level=1)
    p = d.add_paragraph("Лицензиар предоставляет Лицензиату ")
    p.add_run("простую (неисключительную) лицензию").bold = True
    p.add_run(" на использование программы для ЭВМ «Calab» (далее — «Программа»).")
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for item in ("воспроизведение Программы;", "использование по прямому назначению;", "обновление Программы."):
        d.add_paragraph(item, style="List Number")
    d.add_heading("2. Вознаграждение", level=1)
    for item in ("Размер вознаграждения указан в Приложении 1.", "Оплата — в течение 5 рабочих дней."):
        d.add_paragraph(item, style="List Bullet")

    tbl = d.add_table(rows=1, cols=4)
    tbl.style = "Light Grid Accent 1"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, h in zip(tbl.rows[0].cells, ("№", "Тариф", "Срок", "Сумма, ₽")):
        c.text = h
        c.paragraphs[0].runs[0].bold = True
    for row in (("1", "Команда", "12 мес.", "120 000,00"), ("2", "Бизнес", "12 мес.", "360 000,00"), ("", "Итого", "", "480 000,00")):
        cells = tbl.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
    last = tbl.rows[3].cells
    last[0].merge(last[2])
    hdr = OxmlElement("w:tblHeader")  # the header row repeats on every page
    hdr.set(qn("w:val"), "true")
    tbl.rows[0]._tr.get_or_add_trPr().append(hdr)
    red = d.add_paragraph().add_run("Внимание: сумма без НДС (УСН).")
    red.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)

    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_heading("Приложение 1", level=1)
    d.add_paragraph("Реквизиты сторон", style="Heading 2")
    land = d.add_section()
    land.orientation = WD_ORIENT.LANDSCAPE
    land.page_width, land.page_height = Cm(29.7), Cm(21)
    d.add_paragraph("Альбомная страница: график платежей.")
    d.save(path)


def pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor as PRGB
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches, Pt as PPt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
    s = prs.slides.add_slide(prs.slide_layouts[0])  # Title Slide
    s.shapes.title.text = "Итоги квартала"
    s.placeholders[1].text = "Calab · Q3 2026"
    s = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    s.shapes.title.text = "Что сделали"
    tf = s.placeholders[1].text_frame
    tf.text = "Диск и версии файлов"
    for txt, lvl in (("Правка xlsx/docx в окне", 1), ("Телефония", 0), ("Отметки о прочтении", 0)):
        p = tf.add_paragraph()
        p.text, p.level = txt, lvl
    s.notes_slide.notes_text_frame.text = "Заметка докладчика: показать демо."
    s = prs.slides.add_slide(prs.slide_layouts[3])  # Two Content
    s.shapes.title.text = "Было / стало"
    s.placeholders[1].text = "Скачать → править → загрузить"
    s.placeholders[2].text = "Открыть → править"
    s = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only + table + chart + text box
    s.shapes.title.text = "Цифры"
    tbl = s.shapes.add_table(3, 3, Inches(0.5), Inches(1.6), Inches(6), Inches(1.5)).table
    for r, row in enumerate((("Метрика", "Q2", "Q3"), ("Пользователи", "120", "180"), ("Выручка, ₽", "1,2 млн", "1,9 млн"))):
        for c, v in enumerate(row):
            tbl.cell(r, c).text = v
    cd = CategoryChartData()
    cd.categories = ["Июль", "Август", "Сентябрь"]
    cd.add_series("Пользователи", (140, 160, 180))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7), Inches(1.6), Inches(5.5), Inches(4), cd)
    tb = s.shapes.add_textbox(Inches(0.5), Inches(6.3), Inches(6), Inches(0.6))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text, run.font.bold, run.font.size = "Рост +50 %", True, PPt(20)
    run.font.color.rgb = PRGB(0x1F, 0x7A, 0x3A)
    prs.save(path)


def csv_file(path: Path) -> None:
    import csv

    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["ФИО", "Должность", "Начало", "Дней", "Комментарий"])
        w.writerow(["Иванова Анна", "Бухгалтер", "02.07.2026", "14", "с переносом; по согласованию"])
        w.writerow(["Петров Пётр", "Инженер", "15.08.2026", "28", "две части\nиюль и август"])


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "office-fixtures")
    out.mkdir(parents=True, exist_ok=True)
    xlsx(out / "vacations.xlsx")
    docx(out / "contract.docx")
    pptx(out / "deck.pptx")
    csv_file(out / "staff.csv")
    print(out / "vacations.xlsx", out / "contract.docx", out / "deck.pptx", out / "staff.csv")
